"""Lightweight tests for MOJO dish-card PDF import. Run: python -m unittest test_dishes."""
import io
import unittest

import fitz
from PIL import Image
from fastapi import HTTPException
from dish_extract import extract_pdf
from dishes import validate_card,jpg_thumbnail

class DishCards(unittest.TestCase):
    def test_pdf_extracts_image_and_ingredients(self):
        pdf=fitz.open()
        page=pdf.new_page(width=960,height=540)
        page.insert_text((400,40),"SALADS",fontsize=40)
        page.insert_text((40,100),"Caprese salad",fontsize=32)
        for i,name in enumerate(("Tomato","Mozzarella","Basil")):
            page.insert_text((63,190+24*i),name,fontsize=16)
        image=io.BytesIO()
        Image.new("RGB",(500,350),(220,100,40)).save(image,"JPEG")
        page.insert_image(fitz.Rect(550,110,920,350),stream=image.getvalue())
        cards=extract_pdf(pdf.tobytes())
        self.assertEqual(len(cards),1)
        self.assertEqual(cards[0]["title"],"Caprese salad")
        self.assertIn("Mozzarella",cards[0]["ingredients"])
        self.assertGreater(len(cards[0]["photo"]),150)

    def test_structured_powerpoint_template_extracts_fields_and_actual_photo(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx_ttk import template_pptx,extract_pptx
        deck=Presentation(io.BytesIO(template_pptx()))
        slide=deck.slides[1]
        for shape in slide.shapes:
            if not shape.has_text_frame:continue
            for placeholder,actual in [
                ("[Впишите название]","Салат с томатами"),
                ("[Впишите раздел]","Салаты"),
                ("[Ингредиент 1\nИнгредиент 2]","Томаты\nСыр"),
                ("[Впишите описание и подачу]","Подача на холодной тарелке"),
                ("[Перечислите по утверждённой ТТК]","Молоко"),
                ("[Документ и дата]","ТТК версия 2026"),
            ]:
                if placeholder in shape.text:shape.text=shape.text.replace(placeholder,actual)
        photo=io.BytesIO()
        Image.new("RGB",(380,270),(180,80,50)).save(photo,"PNG")
        slide.shapes.add_picture(io.BytesIO(photo.getvalue()),Inches(8.2),Inches(1.2),width=Inches(4))
        data=io.BytesIO();deck.save(data)
        rows=extract_pptx(data.getvalue())
        self.assertEqual(len(rows),1)
        item=rows[0]
        self.assertEqual(item["source_page"],2)
        self.assertEqual(item["title"],"Салат с томатами")
        self.assertEqual(item["ingredients"],["Томаты","Сыр"])
        self.assertEqual(item["allergens"],"Молоко")
        self.assertEqual(item["source_version"],"ТТК версия 2026")
        self.assertTrue(item["photo"])
        self.assertNotIn("Аллергены не подтверждены",item["warning"])

    def test_legacy_full_station_slide_must_not_be_used_as_food_photo(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx_ttk import template_pptx, extract_pptx
        deck=Presentation(io.BytesIO(template_pptx()))
        slide=deck.slides[1]
        for shape in slide.shapes:
            if not shape.has_text_frame:continue
            for old,new in [
                ("[Впишите название]","Уха"),
                ("[Впишите раздел]","Супы"),
                ("[Ингредиент 1\\nИнгредиент 2]","Рыбный бульон\\nКартофель"),
                ("[Впишите описание и подачу]","По утверждённой ТТК"),
                ("[Перечислите по утверждённой ТТК]","Рыба"),
                ("[Документ и дата]","ТТК MOJO v1")
            ]:
                if old in shape.text:shape.text=shape.text.replace(old,new)
        picture=io.BytesIO()
        Image.new("RGB",(1250,703),(230,230,220)).save(picture,"JPEG")
        slide.shapes.add_picture(io.BytesIO(picture.getvalue()),Inches(7),Inches(1),width=Inches(5.5))
        out=io.BytesIO();deck.save(out)
        cards=extract_pptx(out.getvalue())
        self.assertEqual(len(cards),1)
        self.assertEqual(cards[0]["title"],"Уха")
        self.assertEqual(cards[0]["ingredients"],["Рыбный бульон","Картофель"])
        self.assertIsNone(cards[0]["photo"])
        self.assertIn("Не найдено фото",cards[0]["warning"])

    def test_powerpoint_unknown_allergens_never_claims_no_allergens(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx_ttk import extract_pptx
        deck=Presentation()
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        txt=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1))
        txt.text="НАЗВАНИЕ БЛЮДА:\nТестовое блюдо"
        i=io.BytesIO();deck.save(i)
        cards=extract_pptx(i.getvalue())
        self.assertEqual(len(cards),1)
        self.assertEqual(cards[0]["allergens"],"")
        self.assertIn("Аллергены не подтверждены",cards[0]["warning"])

    def test_yandex_maps_only_verified_ingredients_and_allergens(self):
        import json
        import os
        from unittest.mock import patch
        from pptx_ttk import assist_map
        material=[{"page":3,"texts":["Салат с томатами: Томаты, Сыр. Аллергены: молоко."]}]
        suggestion={"cards":[{"page":3,"title":"Салат с томатами",
            "section":"Салаты","ingredients":["Томаты","Сыр","Арахис"],
            "description":"Проверь состав","allergens":"молоко и орехи"}]}
        with patch.dict(os.environ,{"YANDEX_API_KEY":"test-api","YANDEX_FOLDER_ID":"test-folder"}):
            with patch("pptx_ttk.httpx.Client") as client:
                response=client.return_value.__enter__.return_value.post.return_value
                response.json.return_value={"choices":[{"message":{
                    "content":json.dumps(suggestion)}}]}
                mapped=assist_map(material)
        self.assertEqual(mapped[3]["title"],"Салат с томатами")
        self.assertEqual(mapped[3]["ingredients"],["Томаты","Сыр"])
        self.assertEqual(mapped[3]["allergens"],"")
        self.assertEqual(client.return_value.__enter__.return_value.post.call_count,1)

    def test_pdf_ai_can_extract_photo_from_left_of_slide(self):
        from dish_extract import extract_pdf_assisted
        from unittest.mock import patch
        doc=fitz.open(); page=doc.new_page(width=960,height=540)
        pic=io.BytesIO()
        Image.new("RGB",(400,300)).save(pic,"PNG")
        page.insert_image(fitz.Rect(20,40,420,340),stream=pic.getvalue())
        page.insert_text((550,100),"Caprese",fontsize=25)
        page.insert_text((550,180),"Tomato",fontsize=16)
        answer={1:{"title":"Caprese","section":"Salads","ingredients":["Tomato"],
                   "description":"","allergens":""}}
        with patch("pptx_ttk.assist_map",return_value=answer) as mock:
            result=extract_pdf_assisted(doc.tobytes())
        self.assertEqual(len(result),1)
        self.assertTrue(result[0]["photo"])
        self.assertEqual(result[0]["ingredients"],["Tomato"])
        self.assertIn("Аллергены не подтверждены",result[0]["warning"])
        mock.assert_called_once()

    def test_pdf_ai_requires_readable_text_layer(self):
        from dish_extract import extract_pdf_assisted
        doc=fitz.open();page=doc.new_page()
        pic=io.BytesIO();Image.new("RGB",(300,200)).save(pic,"PNG")
        page.insert_image(page.rect,stream=pic.getvalue())
        with self.assertRaises(ValueError):
            extract_pdf_assisted(doc.tobytes())

    def test_reject_invalid_pdf(self):
        with self.assertRaises(ValueError):
            extract_pdf(b"not pdf")

    def test_validate_card(self):
        c=validate_card({"title":"Салат","section":"Салаты","ingredients":["Томаты","Сыр"],"description":"Подача"})
        self.assertEqual(len(c["ingredients"]),2)
        with self.assertRaises(HTTPException):
            validate_card({"title":"","section":"Салаты","ingredients":[],"description":""})

    def test_image_resize(self):
        f=io.BytesIO()
        Image.new("RGB",(4000,1000)).save(f,"PNG")
        data=jpg_thumbnail(f.getvalue())
        with Image.open(io.BytesIO(data)) as im:
            self.assertLessEqual(im.width,900)

@unittest.skipUnless(__import__("os").getenv("DATABASE_URL"),"needs disposable PostgreSQL test schema")
class DishIntegration(unittest.TestCase):
    def setUp(self):
        import os,uuid,psycopg,server
        from pathlib import Path
        from psycopg.rows import dict_row
        from unittest.mock import patch
        from fastapi.testclient import TestClient
        self.schema="dish_test_"+uuid.uuid4().hex
        self.original=server.db
        with self.original() as c:c.execute("CREATE SCHEMA "+self.schema)
        def connect():
            return psycopg.connect(os.environ["DATABASE_URL"],row_factory=dict_row,
                                   options="-c search_path="+self.schema)
        self.connect=connect
        self.patch=patch.object(server,"db",connect)
        self.patch.start()
        with connect() as c:
            c.execute((Path(__file__).parent/"schema.sql").read_text())
            c.execute("""INSERT INTO staff(id,name,role,restaurant,bot_started) VALUES
              (10,'Owner','manager','MOJO',true),
              (11,'Trainee','trainee','MOJO',true),
              (12,'Other','local','OTHER',true)""")
            for uid in (10,11,12):
                c.execute("""INSERT INTO sessions(token_hash,staff_id,expires_at)
                  VALUES(%s,%s,now()+interval '1 hour')""",
                  (__import__("hashlib").sha256(("test-"+str(uid)).encode()).hexdigest(),uid))
        self.client=TestClient(server.app)

    def tearDown(self):
        self.patch.stop()
        with self.original() as c:c.execute("DROP SCHEMA "+self.schema+" CASCADE")

    def headers(self,uid):
        return {"Authorization":"Bearer test-"+str(uid)}

    def sample_pdf(self):
        doc=fitz.open()
        p=doc.new_page(width=960,height=540)
        p.insert_text((400,40),"SALADS",fontsize=40)
        p.insert_text((40,100),"Caprese salad",fontsize=32)
        for i,name in enumerate(("Tomato","Mozzarella","Basil")):
            p.insert_text((63,190+i*24),name,fontsize=16)
        pic=io.BytesIO();Image.new("RGB",(500,350),(180,80,40)).save(pic,"JPEG")
        p.insert_image(fitz.Rect(550,110,920,350),stream=pic.getvalue())
        return doc.tobytes()

    def test_pptx_template_import_requires_manager_and_allergen_review(self):
        from pptx import Presentation
        from pptx_ttk import template_pptx
        from pptx.util import Inches
        deck=Presentation(io.BytesIO(template_pptx()))
        slide=deck.slides[1]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for old,new in [
                    ("[Впишите название]","Тестовое блюдо"),
                    ("[Впишите раздел]","Тестовый раздел"),
                    ("[Ингредиент 1\nИнгредиент 2]","Томаты\nСыр"),
                    ("[Впишите описание и подачу]","Тестовая подача"),
                    ("[Документ и дата]","ТТК v1")
                ]:shape.text=shape.text.replace(old,new)
        picture=io.BytesIO();Image.new("RGB",(180,130)).save(picture,"PNG")
        slide.shapes.add_picture(io.BytesIO(picture.getvalue()),Inches(8.1),Inches(1.0),width=Inches(4))
        output=io.BytesIO();deck.save(output)
        template=self.client.get("/api/dishes/template",headers=self.headers(10))
        self.assertEqual(template.status_code,200,template.text[:100])
        self.assertGreater(len(template.content),1000)
        self.assertEqual(self.client.get("/api/dishes/template",
            headers=self.headers(11)).status_code,403)
        route="/api/dishes/import?filename=MOJO_test.pptx"
        mimetype={"Content-Type":"application/vnd.openxmlformats-officedocument.presentationml.presentation"}
        self.assertEqual(self.client.post(route,content=output.getvalue(),
            headers=self.headers(11)|mimetype).status_code,403)
        result=self.client.post(route,content=output.getvalue(),headers=self.headers(10)|mimetype)
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()["created"],1)
        row=self.client.get("/api/dishes",headers=self.headers(10)).json()[0]
        self.assertFalse(row["published"])
        self.assertEqual(row["title"],"Тестовое блюдо")
        self.assertEqual(row["ingredients"],["Томаты","Сыр"])
        self.assertIn("ТТК v1",row["source_name"])
        self.assertFalse(row["allergens"])
        self.assertEqual(self.client.get("/api/dishes/"+str(row["id"])+"/photo",
            headers=self.headers(10)).status_code,200)
        self.assertEqual(self.client.post(f'/api/dishes/{row["id"]}/publish',
            json={"reviewed":True},headers=self.headers(10)).status_code,409)
        edit={"title":row["title"],"section":row["section"],"ingredients":row["ingredients"],
              "description":"Подача после проверки","allergens":"Молоко"}
        saved=self.client.put(f'/api/dishes/{row["id"]}',headers=self.headers(10),json=edit)
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertEqual(self.client.post(f'/api/dishes/{row["id"]}/publish',
            json={"reviewed":True},headers=self.headers(10)).status_code,200)

    def test_manager_pdf_ai_mode_creates_private_unapproved_card(self):
        from unittest.mock import patch
        doc=fitz.open();page=doc.new_page(width=800,height=500)
        page.insert_text((460,90),"Caprese",fontsize=27)
        page.insert_text((460,180),"Tomato",fontsize=17)
        pic=io.BytesIO();Image.new("RGB",(280,200)).save(pic,"PNG")
        page.insert_image(fitz.Rect(45,110,350,350),stream=pic.getvalue())
        uri="/api/dishes/import?filename=menu.pdf&mode=assist"
        headers=self.headers(10)|{"Content-Type":"application/pdf"}
        self.assertEqual(self.client.post(uri,content=doc.tobytes(),
            headers=self.headers(11)|{"Content-Type":"application/pdf"}).status_code,403)
        answer={1:{"title":"Caprese","section":"Salads","ingredients":["Tomato"],
                   "description":"Present on plate","allergens":""}}
        with patch("pptx_ttk.assist_map",return_value=answer):
            result=self.client.post(uri,content=doc.tobytes(),headers=headers)
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()["format"],"pdf_ai")
        self.assertEqual(result.json()["created"],1)
        self.assertEqual(self.client.get("/api/dishes",
            headers=self.headers(11)).json(),[])
        dish=self.client.get("/api/dishes",headers=self.headers(10)).json()[0]
        self.assertEqual(dish["ingredients"],["Tomato"])
        self.assertFalse(dish["published"])
        self.assertEqual(self.client.post(f'/api/dishes/{dish["id"]}/publish',
            headers=self.headers(10),json={"reviewed":True}).status_code,409)
        edited={"title":"Caprese","section":"Salads","ingredients":["Tomato"],
                "description":"Verified by manager","allergens":"Milk"}
        self.assertEqual(self.client.put(f'/api/dishes/{dish["id"]}',
            headers=self.headers(10),json=edited).status_code,200)
        self.assertEqual(self.client.post(f'/api/dishes/{dish["id"]}/publish',
            headers=self.headers(10),json={"reviewed":True}).status_code,200)

    def test_manager_import_review_photo_and_archive(self):
        data=self.sample_pdf()
        h={**self.headers(10),"Content-Type":"application/pdf"}
        endpoint="/api/dishes/import"
        self.assertEqual(self.client.post(endpoint,content=data,headers=self.headers(11)|{"Content-Type":"application/pdf"}).status_code,403)
        created=self.client.post(endpoint,content=data,headers=h)
        self.assertEqual(created.status_code,200,created.text)
        self.assertEqual(created.json()["created"],1)
        again=self.client.post(endpoint,content=data,headers=h)
        self.assertEqual(again.json()["skipped"],1)
        cards=self.client.get("/api/dishes",headers=self.headers(10)).json()
        self.assertEqual(len(cards),1)
        card=cards[0]
        cid=card["id"]
        self.assertFalse(card["published"])
        self.assertEqual(self.client.get("/api/dishes",headers=self.headers(11)).json(),[])
        self.assertEqual(self.client.get(f"/api/dishes/{cid}/photo",headers=self.headers(11)).status_code,404)
        self.assertEqual(self.client.get(f"/api/dishes/{cid}/photo",headers=self.headers(10)).status_code,200)
        self.assertEqual(self.client.post(f"/api/dishes/{cid}/publish",headers=self.headers(11),json={"reviewed":True}).status_code,403)
        edit={"title":"Caprese salad","section":"SALADS","ingredients":["Tomato","Mozzarella","Basil"],
              "description":"Manager-reviewed synthetic description"}
        self.assertEqual(self.client.put(f"/api/dishes/{cid}",json=edit,headers=self.headers(12)).status_code,404)
        self.assertEqual(self.client.put(f"/api/dishes/{cid}",json=edit,headers=self.headers(10)).status_code,200)
        self.assertEqual(self.client.post(f"/api/dishes/{cid}/publish",json={"reviewed":True},headers=self.headers(10)).status_code,200)
        self.assertEqual(len(self.client.get("/api/dishes",headers=self.headers(11)).json()),1)
        self.assertEqual(self.client.delete(f"/api/dishes/{cid}",headers=self.headers(10)).status_code,200)
        self.assertEqual(self.client.get("/api/dishes",headers=self.headers(11)).json(),[])

    def test_ai_practice_failure_retry_is_idempotent(self):
        import uuid
        from unittest.mock import patch
        from fastapi import HTTPException
        from psycopg.types.json import Jsonb
        with self.connect() as c:
            cid=c.execute("""INSERT INTO courses(author,restaurant,title,kind,body,source,
                questions,pass_score,published)
                VALUES (10,'MOJO','Тестовый курс','ttk','Состав','Синтетический',
                        %s,80,true) RETURNING id""",(Jsonb([]),)).fetchone()["id"]
            did=c.execute("""INSERT INTO dish_cards(restaurant,created_by,title,section,
              ingredients,description,photo,published)
              VALUES('MOJO',10,'Салат','Салаты',%s,'Подаём охлаждённым',%s,true)
              RETURNING id""",(Jsonb(["Томаты","Сыр"]),b"test-photo")).fetchone()["id"]
            c.execute("INSERT INTO course_dishes(course_id,dish_id,position) VALUES(%s,%s,0)",(cid,did))
        uri=f"/api/learning/{cid}/dish-practice"
        rid=str(uuid.uuid4())
        payload={"dish_id":did,"request_id":rid,"answer":"Сначала уточню ограничения гостя, затем расскажу о проверенном составе блюда."}
        feedback={"summary":"Разбор","strengths":"Верно","improvements":"Сверяй ТТК","next_step":"Повторить"}
        with patch.dict(__import__("os").environ,{"YANDEX_API_KEY":"test-only","YANDEX_FOLDER_ID":"test-folder"}):
            with patch("learning.evaluate_with_yandex",side_effect=[HTTPException(502,"No provider"),feedback]) as call:
                first=self.client.post(uri,json=payload,headers=self.headers(11))
                self.assertEqual(first.status_code,502,first.text)
                second=self.client.post(uri,json=payload,headers=self.headers(11))
                self.assertEqual(second.status_code,200,second.text)
                self.assertEqual(second.json(),feedback)
                third=self.client.post(uri,json=payload,headers=self.headers(11))
                self.assertEqual(third.json(),feedback)
                self.assertEqual(call.call_count,2)
        with self.connect() as c:
            row=c.execute("SELECT status,tries FROM dish_practice_attempts WHERE request_id=%s",(rid,)).fetchone()
            self.assertEqual((row["status"],row["tries"]),("done",2))

    def test_curated_quiz_replaces_legacy_and_saves_question_snapshot(self):
        import uuid
        from psycopg.types.json import Jsonb
        from quiz_cards import approved_card_questions
        from dish_lessons import dish_chapters
        old=[dict(type="single",question="Какой компонент прямо указан в станционке?",options=["Томат","Сыр","Укроп"],correct=1,explanation="Сыр",source="старый тест")]
        with self.connect() as c:
            cid=c.execute("""INSERT INTO courses(author,restaurant,title,kind,body,source,questions,pass_score,published)
                VALUES(10,'MOJO','Кухня. Салаты','ttk','Состав блюд','Синтетический пример',%s,80,true)
                RETURNING id""",(Jsonb(old),)).fetchone()["id"]
            for i,(title,ingredients) in enumerate([
                ("Салат с томатами",["Томаты","Сыр","Оливковое масло","Руккола","Базилик","Соль"]),
                ("Салат с огурцами",["Огурцы","Укроп","Сметана","Яйцо","Горчица","Петрушка"]),
            ]):
                did=c.execute("""INSERT INTO dish_cards(restaurant,created_by,title,section,ingredients,description,photo,published)
                    VALUES('MOJO',10,%s,'Салаты',%s,'Утверждённая подача',%s,true)
                    RETURNING id""",(title,Jsonb(ingredients),b'photo')).fetchone()["id"]
                c.execute("INSERT INTO course_dishes(course_id,dish_id,position) VALUES(%s,%s,%s)",(cid,did,i))
            expected=approved_card_questions(dish_chapters(c,cid,'MOJO'),cid)
        course=self.client.get(f"/api/learning/{cid}",headers=self.headers(11))
        self.assertEqual(course.status_code,200,course.text)
        details=course.json()
        kinds={q["type"] for q in details["questions"]}
        self.assertEqual(kinds,{"ingredients","multiple"})
        self.assertNotIn("прямо указан",str(details["questions"]))
        for ch in details["chapters"]:
            result=self.client.post(f"/api/learning/{cid}/read",headers=self.headers(11),json={"chapter_id":ch["id"]})
            self.assertEqual(result.status_code,200,result.text)
        responses=[
            "\n".join(group[0] for group in q["ingredients"]) if q["type"]=="ingredients" else q["correct"]
            for q in expected
        ]
        result=self.client.post(f"/api/learning/{cid}/submit",headers=self.headers(11),json={
            "request_id":str(uuid.uuid4()),"answers":responses
        })
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()["score"],100)
        self.assertEqual(len(result.json()["review"]),len(expected))
        with self.connect() as c:
            stored=c.execute("SELECT questions FROM learning_attempts WHERE course_id=%s",(cid,)).fetchone()
            self.assertEqual(len(stored["questions"]),len(expected))
        # All approved ingredients are valid in free answers, not just the selected subset.
        from questions import check
        first=next(q for q in expected if q["type"]=="ingredients")
        answer="\n".join([g[0] for g in first["ingredients"]]+first["allowed_extra"])
        self.assertEqual(check(first,answer)["fraction"],1)

    def test_voice_route_requires_login_and_does_not_store_audio(self):
        from unittest.mock import patch
        endpoint="/api/voice/transcribe"
        media=b"\x00\x00\x00\x00ftyp" + b"0"*1024
        with patch("voice.normalize_audio",return_value=b"OggS-test"),patch("voice.yandex_recognize",return_value="Я уточню предпочтения гостя.") as provider:
            denied=self.client.post(endpoint,content=media,headers={"Content-Type":"audio/mp4"})
            self.assertEqual(denied.status_code,401)
            response=self.client.post(endpoint,content=media,headers=self.headers(11)|{"Content-Type":"audio/mp4"})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(response.json()["text"],"Я уточню предпочтения гостя.")
            self.assertFalse(response.json()["audio_saved"])
            self.assertEqual(provider.call_count,1)
        with self.connect() as c:
            log=c.execute("SELECT staff_id FROM voice_transcriptions").fetchall()
            self.assertEqual(log,[{"staff_id":11}])

if __name__=="__main__":
    unittest.main()
